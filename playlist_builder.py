# -*- coding: utf-8 -*-
"""
محرك البناء اليومي.

لكل يوم: يقرأ لقطة لوحة "ترتيب 8 ساعات" (مكرّرة لتغطي 24 ساعة) مع بدائل الفجر/المساء/الجمعة،
ويبني تسلسل ذلك اليوم (بادئة + مادة + فواصل لكل فقرة) دون نسخ أي ملف؛ والنسخ خطوة منفصلة.

مبدأ الدقة الزمنية (مؤشر زمني):
  يتتبّع المحرك "المؤشر" = مجموع ما بُني فعلياً من مدد، ويقارنه بالزمن المتوقع لنهاية كل
  فقرة (رقم العمود × مدة الفقرة − ما تُرك فارغاً عمداً). وتُملأ الفواصل بمقدار الفرق بينهما
  بدقة (± دقة الإغلاق بالثواني)، فلا تتراكم الأخطاء أبداً، وتنتهي كل فقرة (30 دقيقة أو 60...)
  وكل 8 ساعات وكل 24 ساعة في موعدها بالضبط ضمن هذا الهامش الصغير.

آلية اختيار الفواصل لملء R ثانية:
  1) بوابة الوجود: لكل نوع احتمال "وجود" في هذه الفقرة (وإن لم يجتزها أي نوع يُستخدم كلها
     حتى لا تبقى الفقرة ناقصة).
  2) اختيار عشوائي لكل ملف ضمن نافذة المدة المقبولة: الهدف = المتبقي × (نسبة تحديد المدة للنوع
     ÷ مجموع النسب)، ضمن هامش التسامح (يتوسّع تدريجياً)، بشرط ألا يتجاوز الملف المتبقي، وأن
     يبقى بعده وقت يمكن إغلاقه (صفر تقريباً أو ≥ أقصر ملف).
  3) عند اقتراب النهاية يُبحث بدقة (bisect) عن ملف واحد أو ملفين يغلقان المتبقي تماماً.
  4) تقليل التكرار: وزن الاختيار يتناقص بشدة لكل ملف سبق استخدامه في اليوم (وبدرجة أقل في
     بقية الأيام)، ولا يتكرر ملف داخل الفقرة الواحدة.
"""
import math
import os
import random
import re
import shutil
from bisect import bisect_left, bisect_right
from collections import Counter
from datetime import datetime, timedelta
from pathlib import Path
from typing import Callable, Dict, List, Optional, Tuple

from audio_utils import get_duration_cached
from library_scanner import IntroIndex, scan_intro_folder
from models import FasilConfig, GridSnapshot, OverrideSnapshot

LogFn = Callable[[str], None]
ProgressFn = Callable[[int, int], None]

DEFAULT_CLOSE_TOLERANCE = 1.0   # ثانية: أقصى انحراف مقبول لإغلاق الفقرة
MAX_ATTEMPTS = 40
CLOSE_ZONE = 300.0              # عند بلوغ المتبقي هذا الحد (ثوانٍ) يُحاوَل الإغلاق الدقيق فوراً
CHECK_ZONE = 240.0              # اختيار يترك متبقياً أقل من هذا يُتحقَّق أنه قابل للإغلاق
WINDOW_STAGES = 6


class BuildResult:
    def __init__(self):
        # كل عنصر: (الترتيب، الدور، المسار الأصلي، المدة بالثواني، اسم معروض)
        self.entries: List[Tuple[int, str, Path, float, str]] = []
        self.total_duration: float = 0.0
        self.expected_total: float = 0.0   # المدة المتوقعة لليوم (24 ساعة ناقص ما تُرك فارغاً عمداً)
        self.warnings: List[str] = []
        # تخطيط اليوم لجدول المواعيد: (العمود، عدد الفترات، السلسلة، رقم المادة) وبدائل (الاسم، العمود، الفترات)
        self.layout: List[Tuple[int, int, object, int]] = []
        self.layout_overrides: List[Tuple[str, int, int]] = []


# ---------------------------------------------------------------------------
class _Entry:
    __slots__ = ("eff", "path", "dur", "type_idx")

    def __init__(self, eff: float, path: Path, dur: float, type_idx: int):
        self.eff = eff          # المدة الفعلية = مدة الملف + مدة بادئة نوعه (إن وُجدت)
        self.path = path
        self.dur = dur
        self.type_idx = type_idx


class _TypePool:
    def __init__(self, cfg: FasilConfig, idx: int, intro_dur: float):
        self.cfg = cfg
        self.idx = idx
        self.intro_path: Optional[Path] = cfg.intro_path if intro_dur > 0 else None
        self.intro_dur = intro_dur
        self.entries: List[_Entry] = sorted(
            (_Entry(d + intro_dur, p, d, idx) for p, d in cfg.fasil_type.files if d and d > 0),
            key=lambda e: e.eff,
        )
        self.effs: List[float] = [e.eff for e in self.entries]
        self.dur_pct = max(float(cfg.duration_percentage), 0.0)
        self.presence = float(cfg.daily_presence_percentage)

    @property
    def min_eff(self) -> float:
        return self.effs[0] if self.effs else float("inf")


class FasilPool:
    """يملأ وقتاً معيّناً (R ثانية) بفواصل بدقة عالية مع عشوائية وتقليل للتكرار."""

    def __init__(self, fasil_configs: List[FasilConfig], tolerance_ratio: float,
                 eps: float = DEFAULT_CLOSE_TOLERANCE, warn: Optional[Callable[[str], None]] = None):
        self.eps = max(float(eps), 0.05)
        self.tol = max(float(tolerance_ratio), 0.0)
        self.types: List[_TypePool] = []
        self.day_usage: Counter = Counter()
        self.all_usage: Counter = Counter()
        self._union_cache: Dict[Tuple[int, ...], Tuple[List[float], List[_Entry]]] = {}
        warn = warn or (lambda m: None)
        for i, fc in enumerate(fasil_configs):
            intro_dur = 0.0
            if fc.intro_path:
                d = get_duration_cached(fc.intro_path)
                if d:
                    intro_dur = d
                else:
                    warn(f"تعذّرت قراءة بادئة نوع الفواصل '{fc.fasil_type.name}' "
                         f"({Path(fc.intro_path).name}) وسيُتجاهل إدراجها.")
            tp = _TypePool(fc, i, intro_dur)
            if tp.entries:
                self.types.append(tp)

    @property
    def has_files(self) -> bool:
        return bool(self.types)

    def new_day(self):
        self.day_usage.clear()

    # ---- أدوات داخلية ----
    def _weight(self, e: _Entry) -> float:
        return max((0.02 ** min(self.day_usage[e.path], 12)) * (0.7 ** min(self.all_usage[e.path], 10)), 1e-12)

    def _eligible(self) -> List[_TypePool]:
        el = [t for t in self.types if (random.random() * 100.0) < t.presence]
        return el if el else list(self.types)   # لا نترك الفقرة ناقصة إن لم يجتز أي نوع البوابة

    def _union(self, types: List[_TypePool]) -> Tuple[List[float], List[_Entry]]:
        key = tuple(t.idx for t in types)
        cached = self._union_cache.get(key)
        if cached is not None:
            return cached
        merged = sorted((e for t in types for e in t.entries), key=lambda e: e.eff)
        res = ([e.eff for e in merged], merged)
        self._union_cache[key] = res
        return res

    def _feasible(self, rem: float, eff: float, min_eff: float) -> bool:
        r2 = rem - eff
        return abs(r2) <= self.eps or r2 >= min_eff

    def _exact_solutions(self, rem: float, union, used: set, limit: int = 30, sample: int = 400):
        """حلول من ملف أو ملفين مجموع مددهما = rem ضمن ± eps (بحث bisect على القائمة المرتبة)."""
        effs, entries = union
        eps = self.eps
        sols: List[tuple] = []
        lo = bisect_left(effs, rem - eps)
        hi = bisect_right(effs, rem + eps)
        for i in range(lo, hi):
            if entries[i].path not in used:
                sols.append((entries[i],))
                if len(sols) >= limit:
                    return sols
        half = bisect_right(effs, rem / 2.0 + eps)
        if half > 0:
            idxs = range(half) if half <= sample else random.sample(range(half), sample)
            for i in idxs:
                a = entries[i]
                if a.path in used:
                    continue
                t = rem - a.eff
                lo2 = bisect_left(effs, t - eps, i)
                hi2 = bisect_right(effs, t + eps, lo2)
                if lo2 < hi2:
                    cands = [k for k in range(lo2, hi2)
                             if k != i and entries[k].path != a.path and entries[k].path not in used]
                    if cands:
                        sols.append((a, entries[random.choice(cands)]))
                        if len(sols) >= limit:
                            break
        return sols

    def _window(self, t: _TypePool, target: float, rem: float, min_eff: float, used: set) -> List[_Entry]:
        """مرشّحو النوع t حول الهدف بنافذة تتوسّع تدريجياً، وتحقق شرط عدم التجاوز وقابلية الإغلاق."""
        hi_cap = rem + self.eps
        for stage in range(WINDOW_STAGES):
            ratio = self.tol + 0.25 * stage
            lo = target * (1.0 - ratio)
            hi = min(target * (1.0 + ratio), hi_cap)
            i = bisect_left(t.effs, lo)
            j = bisect_right(t.effs, hi)
            c = [e for e in t.entries[i:j] if e.path not in used and self._feasible(rem, e.eff, min_eff)]
            if c:
                return c
        j = bisect_right(t.effs, hi_cap)
        return [e for e in t.entries[:j] if e.path not in used and self._feasible(rem, e.eff, min_eff)]

    def _attempt(self, R: float, types: List[_TypePool], union) -> Optional[List[_Entry]]:
        eps = self.eps
        rem = R
        picks: List[_Entry] = []
        used: set = set()
        min_eff = min(t.min_eff for t in types)
        pct_sum = sum(t.dur_pct for t in types)
        turn = 0
        for _ in range(80):
            if rem <= eps:
                break
            if rem < min_eff - eps:
                return None
            if rem <= CLOSE_ZONE:
                sols = self._exact_solutions(rem, union, used)
                if sols:
                    weights = [math.prod(self._weight(e) for e in s) for s in sols]
                    chosen = random.choices(sols, weights=weights, k=1)[0]
                    picks.extend(chosen)
                    rem -= sum(e.eff for e in chosen)
                    break
            pick = None
            for k in range(len(types)):
                t = types[(turn + k) % len(types)]
                share = (t.dur_pct / pct_sum) if pct_sum > 0 else 1.0 / len(types)
                target = max(rem * share, min_eff)
                cands = self._window(t, target, rem, min_eff, used)
                weights = [self._weight(e) for e in cands]
                tries = 0
                while cands and tries < 15:
                    idx = random.choices(range(len(cands)), weights=weights, k=1)[0]
                    e = cands[idx]
                    r2 = rem - e.eff
                    if (abs(r2) <= eps or r2 >= CHECK_ZONE
                            or self._exact_solutions(r2, union, used | {e.path}, limit=1)):
                        pick = e
                        break
                    del cands[idx]
                    del weights[idx]
                    tries += 1
                if pick is not None:
                    break
            if pick is None:
                return None
            picks.append(pick)
            used.add(pick.path)
            rem -= pick.eff
            turn += 1
        return picks if abs(rem) <= eps else None

    def _fallback(self, R: float, union) -> List[_Entry]:
        """أفضل جهد (نادراً): يملأ من الأسفل بأكبر ملفات لا تتجاوز المتبقي. لا يتجاوز R أبداً."""
        effs, entries = union
        rem, picks, used = R, [], set()
        min_eff = effs[0] if effs else float("inf")
        while rem >= min_eff - self.eps and len(picks) < 60:
            k = bisect_right(effs, rem + self.eps) - 1
            while k >= 0 and entries[k].path in used:
                k -= 1
            if k < 0:
                break
            picks.append(entries[k])
            used.add(entries[k].path)
            rem -= entries[k].eff
        return picks

    # ---- الواجهة العامة ----
    def fill(self, R: float, order: int, result: BuildResult) -> Tuple[int, float]:
        """يملأ R ثانية بفواصل ويضيفها إلى result. يعيد (الترتيب الجديد، مجموع ما أُضيف فعلياً)."""
        if R <= self.eps or not self.types:
            return order, 0.0
        gated = self._eligible()
        all_types = list(self.types)
        picks: Optional[List[_Entry]] = None
        union = self._union(gated)
        # المرحلة 1: الأنواع التي اجتازت بوابة الوجود. المرحلة 2: كل الأنواع (لا نترك الفقرة ناقصة).
        stages = [(gated, MAX_ATTEMPTS // 2)]
        if len(gated) != len(all_types):
            stages.append((all_types, MAX_ATTEMPTS))
        for types, tries in stages:
            union = self._union(types)
            min_eff = union[0][0] if union[0] else float("inf")
            if R < min_eff - self.eps:
                continue      # أقصر فاصل في هذه المرحلة أطول من المتبقي
            for _ in range(tries):
                picks = self._attempt(R, types, union)
                if picks:
                    break
            if picks:
                break
        if not picks:
            union = self._union(all_types)
            if R < (union[0][0] if union[0] else float("inf")) - self.eps:
                return order, 0.0     # يُرحَّل الفرق للفقرة التالية
        if not picks:
            picks = self._fallback(R, union)
        added = 0.0
        for e in picks:
            t = self.types[self._type_pos(e.type_idx)]
            if t.intro_path:
                order += 1
                result.entries.append((order, "بادئة فاصل", t.intro_path, t.intro_dur, t.intro_path.name))
            order += 1
            result.entries.append((order, f"فاصل-{t.cfg.fasil_type.name}", e.path, e.dur, e.path.name))
            added += e.eff
            self.day_usage[e.path] += 1
            self.all_usage[e.path] += 1
        return order, added

    def _type_pos(self, type_idx: int) -> int:
        for pos, t in enumerate(self.types):
            if t.idx == type_idx:
                return pos
        return 0


# ---------------------------------------------------------------------------
class _Override:
    def __init__(self, start: int, span: int, label: str, items: List[Tuple[str, Path, float, str]]):
        self.start = start
        self.span = span
        self.label = label
        self.items = items


class DailyScheduleBuilder:
    """
    محرك بناء يومي. grid: GridSnapshot ، overrides: OverrideSnapshot (لقطات ثابتة).
    لا يعتمد على أي حالة محفوظة؛ كل يوم يُحسب مباشرة من اللقطتين.
    """

    def __init__(
        self,
        grid: GridSnapshot,
        overrides: OverrideSnapshot,
        intro_folder: Optional[Path],
        intro_index: Optional[IntroIndex],
        fasil_configs: List[FasilConfig],
        tolerance_ratio: float,
        close_tolerance: float = DEFAULT_CLOSE_TOLERANCE,
        fill_gaps: bool = False,
        log: Optional[LogFn] = None,
    ):
        self.grid = grid
        self.overrides = overrides
        self.intro_folder = Path(intro_folder) if intro_folder else None
        self.intro_index = intro_index
        self.fasil_configs = fasil_configs
        self.fill_gaps = fill_gaps
        self.tolerance_ratio = tolerance_ratio
        self.close_tolerance = close_tolerance
        self.log = log or (lambda msg: None)
        self._pool: Optional[FasilPool] = None
        self._warned: set = set()
        self._intro_cache: Dict[str, Optional[Tuple[Path, float]]] = {}

    # ------------------------------------------------------------------
    def _warn_once(self, result: BuildResult, key: str, msg: str):
        if key not in self._warned:
            self._warned.add(key)
            result.warnings.append(msg)

    def _get_pool(self, result: BuildResult) -> FasilPool:
        if self._pool is None:
            self._pool = FasilPool(
                self.fasil_configs, self.tolerance_ratio, self.close_tolerance,
                warn=lambda m: self._warn_once(result, "pool:" + m, "تحذير: " + m),
            )
        return self._pool

    def _ensure_intro_index(self):
        if self.intro_index is None or (len(self.intro_index) == 0 and self.intro_folder):
            if self.intro_folder and self.intro_folder.exists():
                self.log("جارٍ فحص مجلد البادئات...")
                self.intro_index = scan_intro_folder(self.intro_folder)
            elif self.intro_index is None:
                self.intro_index = IntroIndex()

    def _intro_for(self, series, result: BuildResult) -> Optional[Tuple[Path, float]]:
        key = str(series.folder)
        if key in self._intro_cache:
            return self._intro_cache[key]
        entry, warn = self.intro_index.lookup(series.series_number, series.type_code) if self.intro_index else (None, None)
        if warn:
            self._warn_once(result, "intro:" + key, f"{series.display_name}: {warn}")
        self._intro_cache[key] = entry
        return entry

    # ------------------------------------------------------------------
    def build_day(self, day) -> BuildResult:
        result = BuildResult()
        g = self.grid
        ov = self.overrides
        self._ensure_intro_index()
        pool = self._get_pool(result)
        pool.new_day()

        row = (day - g.start_date).days
        slot_minutes = max(g.slot_minutes, 1)
        slot_secs = slot_minutes * 60
        cols = (24 * 60) // slot_minutes
        base_cols = max(g.cols, 1)
        eps = pool.eps

        if row < 0 or row >= g.rows:
            result.warnings.append(
                f"تنبيه: التاريخ {day.isoformat()} خارج مدى الأيام المحدد في لوحة ترتيب 8 ساعات "
                "(لن تظهر فيه مواد، فقط البدائل والفواصل)."
            )

        if self.intro_index is not None and len(self.intro_index) == 0:
            self._warn_once(result, "intro-empty",
                            "تحذير: لا توجد بادئات (مجلد البادئات فارغ أو لا يحتوي ملفات مطابقة للنمط).")
        if self.fasil_configs and not pool.has_files:
            self._warn_once(result, "no-fasil", "تحذير: لا يوجد أي نوع فواصل يحتوي ملفات صالحة، لن تُضاف فواصل.")

        # ---- بناء قائمة البدائل (تحلّ محل ما هو مجدول في موعدها) ----
        def col_of(minutes: int) -> int:
            return round(((minutes - g.day_start_minutes) % (24 * 60)) / slot_minutes) % cols

        overrides: List[_Override] = []

        def add_override(label: str, minutes: int, items: List[Tuple[str, Path, float, str]]):
            total = sum(d for _r, _p, d, _n in items)
            if total <= 0:
                result.warnings.append(f"تحذير: {label}: تعذّرت قراءة مدة الملف/الملفات، لن يُدرَج هذا البديل.")
                return
            start = col_of(minutes)
            span = max(1, math.ceil(total / slot_secs - 1e-9))
            span = min(span, cols - start)
            overrides.append(_Override(start, span, label, items))

        if ov.morning_file and ov.morning_duration:
            add_override("بديل الفجر", ov.morning_minutes,
                         [("بديل", ov.morning_file, ov.morning_duration, f"بديل الفجر - {ov.morning_file.name}")])
        elif ov.morning_file:
            result.warnings.append("تحذير: بديل الفجر: تعذّرت قراءة مدة الملف، لن يُدرَج.")
        if ov.evening_file and ov.evening_duration:
            add_override("بديل المساء", ov.evening_minutes,
                         [("بديل", ov.evening_file, ov.evening_duration, f"بديل المساء - {ov.evening_file.name}")])
        elif ov.evening_file:
            result.warnings.append("تحذير: بديل المساء: تعذّرت قراءة مدة الملف، لن يُدرَج.")
        if day.weekday() == 4 and ov.friday_files:
            good = [(p, d) for p, d in ov.friday_files if d]
            if len(good) != len(ov.friday_files):
                result.warnings.append("تحذير: بعض ملفات الجمعة تعذّرت قراءة مدتها وتم تجاهلها.")
            if good:
                add_override("بدائل الجمعة", ov.friday_minutes,
                             [("بديل", p, d, f"بديل الجمعة - {p.name}") for p, d in good])

        ov_start: Dict[int, _Override] = {}
        ov_cover: Dict[int, _Override] = {}
        for o in sorted(overrides, key=lambda x: x.start):
            if any(c in ov_cover for c in range(o.start, o.start + o.span)):
                result.warnings.append(f"تحذير: {o.label} يتداخل مع بديل آخر في الوقت، تم تجاهله.")
                continue
            ov_start[o.start] = o
            result.layout_overrides.append((o.label, o.start, o.span))
            for c in range(o.start, o.start + o.span):
                ov_cover[c] = o

        # ---- المرور على أعمدة اليوم ----
        order = 0
        cursor = 0.0      # مجموع ما بُني فعلياً (ثوانٍ)
        skipped = 0.0     # وقت تُرك فارغاً عمداً (فراغات بلا ملء)
        col = 0
        no_fasil_gap = 0.0
        pieces = g.pieces

        def fill_to(end_col: int):
            nonlocal order, cursor, no_fasil_gap
            R = end_col * slot_secs - skipped - cursor
            if R <= eps:
                return
            if pool.has_files:
                order, added = pool.fill(R, order, result)
                cursor += added
            else:
                no_fasil_gap += R

        def handle_gap(c: int):
            nonlocal skipped
            if self.fill_gaps and pool.has_files:
                fill_to(c + 1)
            else:
                skipped += slot_secs

        while col < cols:
            o = ov_start.get(col)
            if o is not None:
                for role, path, dur, display in o.items:
                    order += 1
                    result.entries.append((order, role, path, dur, display))
                    cursor += dur
                fill_to(col + o.span)
                col += o.span
                continue

            logical_col = col % base_cols
            pid = g.occupied.get((row, logical_col))
            piece = pieces.get(pid) if pid else None
            lesson = None
            if piece is not None and piece.col == logical_col:
                idx = row - piece.start_row
                if 0 <= idx < len(piece.series.materials):
                    mat = piece.series.materials[idx]
                    if mat.duration_seconds:
                        lesson = (piece, idx, mat)
                    else:
                        result.warnings.append(
                            f"تحذير: مادة بلا مدة معروفة تم تخطيها: {piece.series.display_name} - مادة {mat.number}"
                        )
            if lesson is None:
                handle_gap(col)
                col += 1
                continue

            piece, idx, mat = lesson
            display = f"{piece.series.display_name} - مادة {mat.number}"
            intro_entry = self._intro_for(piece.series, result)
            intro_dur = intro_entry[1] if intro_entry else 0.0
            need = mat.duration_seconds + intro_dur
            reserved = piece.row_widths[idx] if idx < len(piece.row_widths) else 1
            needed_slots = max(1, math.ceil(need / slot_secs - 1e-9))
            n_slots = max(reserved, needed_slots)
            if needed_slots > reserved:
                result.warnings.append(
                    f"تحذير: {display}: مدتها مع البادئة أكبر من المساحة المحجوزة في اللوحة "
                    f"({reserved} فترة) فاستُخدمت {needed_slots}؛ قد تُزاح المادة التالية."
                )

            conflict = next((c for c in range(col, col + n_slots) if c in ov_cover), None)
            if conflict is not None:
                result.warnings.append(
                    f"تنبيه: {display} تتعارض مع {ov_cover[conflict].label} فتم تجاوزها في هذا اليوم "
                    "(البديل يحلّ محل ما هو مجدول في موعده)."
                )
                handle_gap(col)
                col += 1
                continue

            result.layout.append((col, n_slots, piece.series, mat.number))
            if intro_entry:
                order += 1
                result.entries.append((order, "بادئة", intro_entry[0], intro_dur, intro_entry[0].name))
                cursor += intro_dur
            order += 1
            result.entries.append((order, "درس", mat.path, mat.duration_seconds, display))
            cursor += mat.duration_seconds
            fill_to(col + n_slots)
            col += n_slots

        result.total_duration = cursor
        result.expected_total = col * slot_secs - skipped
        if no_fasil_gap > eps:
            result.warnings.append(
                f"تنبيه: بقي {no_fasil_gap/60:.1f} دقيقة غير مملوءة لعدم وجود فواصل."
            )
        elif abs(result.total_duration - result.expected_total) > max(2 * eps, 3.0):
            result.warnings.append(
                f"تنبيه: انحراف زمني {result.total_duration - result.expected_total:+.1f} ثانية "
                "عن المتوقع (تعذّر الإغلاق الدقيق، غالباً لقصر الفراغ عن أقصر فاصل متاح)."
            )
        return result

    def build_range(self, start_date, end_date) -> "Dict":
        results = {}
        d = start_date
        while d <= end_date:
            self.log(f"جارٍ بناء يوم {d.isoformat()}...")
            results[d] = self.build_day(d)
            d += timedelta(days=1)
        return results

    # ------------------------------------------------------------------
    _GENERATED_RE = re.compile(r"^\d{4}_[^/\\]+$")

    @staticmethod
    def _clear_generated(folder: Path) -> int:
        """يحذف من مجلد يوم سابق الملفات التي أنشأها البرنامج فقط (0001_درس.mp3 ...) وأي .part."""
        n = 0
        for p in folder.iterdir():
            try:
                if p.is_file() and (DailyScheduleBuilder._GENERATED_RE.match(p.name) or p.name.endswith(".part")):
                    p.unlink()
                    n += 1
            except OSError:
                pass
        return n

    @staticmethod
    def day_folder_has_content(output_root: Path, day) -> bool:
        f = Path(output_root) / day.isoformat()
        try:
            return f.is_dir() and any(f.iterdir())
        except OSError:
            return False

    def copy_day_to_output(
        self, day, result: BuildResult, output_root: Path,
        progress: Optional[ProgressFn] = None,
        replace_existing: bool = True,
        cancel: Optional[Callable[[], bool]] = None,
    ) -> Tuple[Path, List[Tuple[int, str, str, str, float]], List[Tuple[str, str]]]:
        """
        ينسخ ملفات يوم واحد إلى مجلد فرعي باسم التاريخ داخل مجلد الإخراج. كل ملف يُنسخ إلى
        اسم مؤقت (.part) ثم يُعاد تسميته، فلا يبقى ملف ناقص عند الانقطاع. فشل ملف لا يوقف
        الباقي (يُسجَّل في الأخطاء). يعيد (مجلد اليوم، صفوف الجدول، قائمة الأخطاء [(المصدر، السبب)]).
        """
        day_folder = Path(output_root) / day.isoformat()
        day_folder.mkdir(parents=True, exist_ok=True)
        if replace_existing:
            self._clear_generated(day_folder)

        rows: List[Tuple[int, str, str, str, float]] = []
        errors: List[Tuple[str, str]] = []
        total = len(result.entries)
        for i, (order, role, src, dur, _display) in enumerate(result.entries, start=1):
            if cancel and cancel():
                break
            safe_role = role.split("-")[0].replace(" ", "_")
            dest_name = f"{order:04d}_{safe_role}{Path(src).suffix}"
            dest_path = day_folder / dest_name
            tmp_path = day_folder / (dest_name + ".part")
            try:
                shutil.copy2(src, tmp_path)
                os.replace(tmp_path, dest_path)
                rows.append((order, role, str(src), dest_name, dur))
                self.log(f"[{day.isoformat()} {i}/{total}] {Path(src).name}  ->  {dest_name}")
            except Exception as e:  # noqa: BLE001
                try:
                    if tmp_path.exists():
                        tmp_path.unlink()
                except OSError:
                    pass
                errors.append((str(src), str(e)))
                rows.append((order, role, str(src), "⚠ فشل النسخ: " + str(e)[:80], dur))
                self.log(f"⚠ [{day.isoformat()} {i}/{total}] فشل نسخ {Path(src).name}: {e}")
            if progress:
                progress(i, total)

        return day_folder, rows, errors

    # ------------------------------------------------------------------
    # صفحة "جدول المواعيد": نفس محتوى لوحة 24 ساعة (صف لكل يوم) + رقم الحلقة + نص مواعيد السلاسل
    # ------------------------------------------------------------------
    WEEKDAYS_AR = ["الاثنين", "الثلاثاء", "الأربعاء", "الخميس", "الجمعة", "السبت", "الأحد"]
    SHEET_TITLE = "جدول المواعيد"
    _PALETTE = ["4C6EF5", "12B886", "F59F00", "E8590C", "E64980", "7048E8", "1098AD", "82C91E",
                "F03E3E", "845EF7", "0CA678", "FAB005", "3B5BDB", "0B7285", "D9480F", "A61E4D"]
    _OVERRIDE_STYLE = {"بديل الفجر": ("صباح", "F59F00", "فترة الصباح"),
                       "بديل المساء": ("مساء", "7048E8", "فترة المساء"),
                       "بدائل الجمعة": ("الجمعة", "12B886", "فترة الجمعة")}

    @staticmethod
    def fmt_hm(total_minutes: int) -> str:
        total_minutes %= 24 * 60
        return f"{total_minutes // 60:02d}:{total_minutes % 60:02d}"

    @staticmethod
    def fmt_ar_time(total_minutes: int) -> str:
        """6:00 صباحاً / 12:00 ظهراً / 8:00 مساءاً"""
        total_minutes %= 24 * 60
        h24, m = divmod(total_minutes, 60)
        h12 = h24 % 12 or 12
        period = "صباحاً" if h24 < 12 else ("ظهراً" if h24 < 15 else "مساءاً")
        return f"{h12}:{m:02d} {period}"

    def schedule_text(self, day, result: "BuildResult") -> str:
        """نص مواعيد سلاسل اليوم: سطر لكل سلسلة بأوقات بثها (مرة كل 8 ساعات مثلاً) ثم أوقات البدائل."""
        g = self.grid
        slot = max(g.slot_minutes, 1)
        base_cols = max(g.cols, 1)

        def t(col: int) -> str:
            return self.fmt_hm(g.day_start_minutes + col * slot)

        groups: Dict[int, List] = {}
        for col, _span, series, _mat in sorted(result.layout, key=lambda x: x[0]):
            groups.setdefault(col % base_cols, []).append((col, series))
        lines = [f"جدول مواعيد السلاسل بتاريخ {day.isoformat()}"]
        for logical in sorted(groups):
            items = groups[logical]
            times = " ".join(t(c) for c, _s in items)
            lines.append(f"{times} {items[0][1].series_name or items[0][1].series_number}")
        for label, start, _span in sorted(result.layout_overrides, key=lambda x: x[1]):
            title = self._OVERRIDE_STYLE.get(label, (label, "", label))[2]
            lines.append(f"{title} الساعة {self.fmt_ar_time(g.day_start_minutes + start * slot)}")
        return "\n".join(lines)

    def write_schedule_sheet(self, wb, results: Dict) -> None:
        """يضيف (أو يستبدل) صفحة جدول المواعيد في مصنف openpyxl لأيام results."""
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        g = self.grid
        slot = max(g.slot_minutes, 1)
        cols = (24 * 60) // slot
        if self.SHEET_TITLE in wb.sheetnames:
            del wb[self.SHEET_TITLE]
        ws = wb.create_sheet(title=self.SHEET_TITLE, index=0)
        ws.sheet_view.rightToLeft = True

        head_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
        head_font = Font(bold=True, color="FFFFFF")
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        headers = ["رقم الحلقة", "التاريخ", "اليوم"] + [self.fmt_hm(g.day_start_minutes + c * slot) for c in range(cols)]
        text_col = len(headers) + 1
        headers.append("مواعيد السلاسل")
        for i, title in enumerate(headers, start=1):
            c = ws.cell(row=1, column=i, value=title)
            c.fill, c.font, c.alignment = head_fill, head_font, center
        ws.row_dimensions[1].height = 22

        for r, day in enumerate(sorted(results), start=2):
            res = results[day]
            episode = (day - g.start_date).days + 1
            ws.cell(row=r, column=1, value=episode).alignment = center
            ws.cell(row=r, column=2, value=day.isoformat()).alignment = center
            ws.cell(row=r, column=3, value=self.WEEKDAYS_AR[day.weekday()]).alignment = center
            for col, span, series, mat_no in res.layout:
                c0 = 4 + col
                span = max(1, min(span, cols - col))
                try:
                    color = self._PALETTE[int(series.series_number) % len(self._PALETTE)]
                except (TypeError, ValueError):
                    color = self._PALETTE[0]
                cell = ws.cell(row=r, column=c0, value=f"{series.series_number} - {mat_no}")
                cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
                cell.font = Font(bold=True, color="FFFFFF", size=9)
                cell.alignment = center
                if span > 1:
                    ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + span - 1)
            for label, start, span in res.layout_overrides:
                word, color, _title = self._OVERRIDE_STYLE.get(label, (label, "808080", label))
                span = max(1, min(span, cols - start))
                cell = ws.cell(row=r, column=4 + start, value=word)
                cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
                cell.font = Font(bold=True, color="FFFFFF", size=9)
                cell.alignment = center
                if span > 1:
                    ws.merge_cells(start_row=r, start_column=4 + start, end_row=r, end_column=4 + start + span - 1)
            text = self.schedule_text(day, res)
            tc = ws.cell(row=r, column=text_col, value=text)
            tc.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
            ws.row_dimensions[r].height = min(409, max(30, 15 * (text.count("\n") + 1)))

        ws.column_dimensions["A"].width = 11
        ws.column_dimensions["B"].width = 13
        ws.column_dimensions["C"].width = 11
        for c in range(cols):
            ws.column_dimensions[get_column_letter(4 + c)].width = 9
        ws.column_dimensions[get_column_letter(text_col)].width = 62
        ws.freeze_panes = "D2"

    def export_schedule_workbook(self, results: Dict, path: Path) -> Path:
        """يحفظ جدول المواعيد وحده في ملف xlsx (دون نسخ أي ملف صوتي)."""
        import openpyxl
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        self.write_schedule_sheet(wb, results)
        wb.save(path)
        return Path(path)

    def write_combined_schedule_workbook(self, day_rows: Dict, output_root: Path, results: Optional[Dict] = None) -> Path:
        """
        يبني/يحدّث جدولاً واحداً (Excel) فيه صفحة لكل يوم، في جذر مجلد الإخراج. إن وُجد الملف
        من تخريج سابق تُستبدل صفحات الأيام المخرَّجة الآن فقط وتبقى بقية الصفحات. وإن كان الملف
        مفتوحاً في Excel (يمنع الحفظ) يُحفظ باسم بديل يحمل الوقت.
        """
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill
        from openpyxl.utils import get_column_letter

        out_path = Path(output_root) / "جدول_التخريج.xlsx"
        wb = None
        if out_path.exists():
            try:
                wb = openpyxl.load_workbook(out_path)
            except Exception:  # noqa: BLE001
                wb = None
        if wb is None:
            wb = openpyxl.Workbook()
            wb.remove(wb.active)

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
        headers = ["الترتيب", "الدور", "الملف الأصلي", "اسم الملف الجديد", "المدة (ثانية)"]

        for day in sorted(day_rows.keys()):
            sheet_name = day.isoformat()[:31]
            if sheet_name in wb.sheetnames:
                del wb[sheet_name]
            ws = wb.create_sheet(title=sheet_name)
            ws.sheet_view.rightToLeft = True
            for col_idx, title in enumerate(headers, start=1):
                cell = ws.cell(row=1, column=col_idx, value=title)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for row_idx, (order, role, src, dest_name, dur) in enumerate(day_rows[day], start=2):
                ws.cell(row=row_idx, column=1, value=order)
                ws.cell(row=row_idx, column=2, value=role)
                ws.cell(row=row_idx, column=3, value=src)
                ws.cell(row=row_idx, column=4, value=dest_name)
                ws.cell(row=row_idx, column=5, value=round(dur, 1))
            for i, w in enumerate([10, 18, 70, 30, 14], start=1):
                ws.column_dimensions[get_column_letter(i)].width = w
            ws.freeze_panes = "A2"

        wb._sheets.sort(key=lambda s: s.title)
        if results:
            self.write_schedule_sheet(wb, {d: r for d, r in results.items() if d in day_rows})
        try:
            wb.save(out_path)
            return out_path
        except PermissionError:
            alt = Path(output_root) / f"جدول_التخريج_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
            wb.save(alt)
            self.log(f"⚠ تعذّر الكتابة على {out_path.name} (ربما مفتوح في Excel)، حُفظ باسم: {alt.name}")
            return alt
